from openpyxl import load_workbook, Workbook


# noinspection SpellCheckingInspection
from openpyxl.worksheet.table import Table, TableStyleInfo


class ExcelGenerator:
    @staticmethod
    def generateExcel(mvfid, startdate, estimatedenddate, actualenddate, noofmembers, unplannedabsences, mvfdata):
        wb = load_workbook('MVF_breakdown.xltx')
        wb.template = False

        data = wb['Data']

        data.append(["Type", "Subtype", "MVFID", "Story Points", "Actual Days"])
        for row in mvfdata:
            data.append(row)

        tab = Table(displayName="Table1", ref="A1:E21")

        data.add_table(tab)

        mvfbreakdown = wb['MVF Breakdown']
        mvfbreakdown["D5"].value = startdate
        mvfbreakdown["D6"].value = actualenddate
        mvfbreakdown["D7"].value = estimatedenddate
        mvfbreakdown["D11"].value = noofmembers
        mvfbreakdown["D12"].value = unplannedabsences

        wb.save(mvfid + '.xlsx')

        return True
